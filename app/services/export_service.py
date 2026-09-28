"""
Export Service generating real Excel (.xlsx) and PDF (.pdf) documents.
Uses OpenPyXL for styled spreadsheets and ReportLab for professional printable PDF reports.
"""

import io
from datetime import datetime, date
from typing import List, Dict, Any

# OpenPyXL imports
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# ReportLab imports
from reportlab.lib.pagesizes import letter, landscape, A4
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle

class ExportService:

    @staticmethod
    def _apply_excel_styling(ws, title: str):
        """Applies professional corporate styling to worksheet headers and borders."""
        # Title banner
        ws.merge_cells("A1:H1")
        title_cell = ws["A1"]
        title_cell.value = f"CHEMICAL STORE MANAGEMENT SYSTEM - {title.upper()}"
        title_cell.font = Font(name="Arial", size=13, bold=True, color="FFFFFF")
        title_cell.fill = PatternFill(start_color="0C4A6E", end_color="0C4A6E", fill_type="solid")
        title_cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[1].height = 28

        # Subtitle timestamp
        ws.merge_cells("A2:H2")
        sub_cell = ws["A2"]
        sub_cell.value = f"Generated On: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')} | Confidential Store Document"
        sub_cell.font = Font(name="Arial", size=9, italic=True, color="64748B")
        sub_cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[2].height = 18

    @staticmethod
    def _format_columns_and_borders(ws, header_row_idx: int):
        """Formats headers, data cells, thin borders, and auto-fits column widths."""
        header_fill = PatternFill(start_color="0284C7", end_color="0284C7", fill_type="solid")
        header_font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        border_thin = Border(
            left=Side(style='thin', color='CBD5E1'),
            right=Side(style='thin', color='CBD5E1'),
            top=Side(style='thin', color='CBD5E1'),
            bottom=Side(style='thin', color='CBD5E1')
        )

        ws.row_dimensions[header_row_idx].height = 24
        for cell in ws[header_row_idx]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = border_thin

        for row in ws.iter_rows(min_row=header_row_idx + 1, max_row=ws.max_row):
            ws.row_dimensions[row[0].row].height = 20
            for cell in row:
                cell.border = border_thin
                cell.font = Font(name="Arial", size=9)
                if isinstance(cell.value, (int, float)):
                    cell.number_format = "#,##0.00"
                    cell.alignment = Alignment(horizontal="right", vertical="center")
                elif isinstance(cell.value, (datetime, date)):
                    cell.number_format = "YYYY-MM-DD"
                    cell.alignment = Alignment(horizontal="center", vertical="center")

        # Auto column width
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val_str = str(cell.value or '')
                if cell.row in (1, 2):
                    continue
                if len(val_str) > max_len:
                    max_len = len(val_str)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    @classmethod
    def export_stock_balance_excel(cls, rows: List[Dict[str, Any]]) -> bytes:
        """Generates Excel workbook for Stock Balance Report."""
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Stock Balance"

        cls._apply_excel_styling(ws, "Stock Balance by Chemical & Batch")

        headers = [
            "Chemical Name", "Code", "Batch Number", "Mfg Date", "Expiry Date",
            "Storage Location", "Supplier", "Received Qty", "Current Stock", "Unit", "Status"
        ]
        ws.append([]) # Empty row 3
        ws.append(headers) # Row 4

        for r in rows:
            ws.append([
                r["chemical_name"], r["chemical_code"], r["batch_number"],
                r["manufacturing_date"] or "N/A", r["expiry_date"],
                r["storage_location_name"], r.get("supplier_name") or "—",
                r["received_quantity"], r["current_quantity"], r["unit"], r["status"]
            ])

        cls._format_columns_and_borders(ws, 4)

        out = io.BytesIO()
        wb.save(out)
        return out.getvalue()

    @classmethod
    def export_daily_report_excel(cls, date_str: str, rows: List[Dict[str, Any]]) -> bytes:
        """Generates Excel workbook for Daily Store Movement Report."""
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Daily Movement"

        cls._apply_excel_styling(ws, f"Daily Store Movement ({date_str})")

        headers = [
            "Chemical", "Batch Number", "Expiry Date", "Storage Location",
            "Opening Stock", "Received", "Used", "Transfer In", "Transfer Out",
            "Production In", "Production Consumed", "Adjustments", "Closing Stock", "Unit"
        ]
        ws.append([])
        ws.append(headers)

        for r in rows:
            ws.append([
                r["chemical_name"], r["batch_number"], r["expiry_date"], r["storage_location"],
                r["opening_stock"], r["received"], r["used"], r["transfer_in"], r["transfer_out"],
                r["production_in"], r["production_consumption"], r["adjustments"], r["closing_stock"], r["unit"]
            ])

        cls._format_columns_and_borders(ws, 4)

        out = io.BytesIO()
        wb.save(out)
        return out.getvalue()

    @classmethod
    def export_ammonia_audit_excel(cls, audit_data: Dict[str, Any]) -> bytes:
        """Generates multi-sheet Excel for Ammonia Barrels, Barrel History, and 10% Preparations."""
        wb = openpyxl.Workbook()
        
        # Sheet 1: 25% Ammonia Barrels
        ws1 = wb.active
        ws1.title = "25% Barrels Ledger"
        cls._apply_excel_styling(ws1, "25% Ammonia Barrel Stock")
        headers1 = ["Barrel ID", "Batch Number", "Original Qty (kg)", "Current Qty (kg)", "Supplier", "Received Date", "Expiry Date", "Storage Location", "Status"]
        ws1.append([])
        ws1.append(headers1)
        for b in audit_data.get("barrels", []):
            ws1.append([
                b["barrel_id"], b["batch_number"], b["original_quantity"], b["current_quantity"],
                b["supplier_name"] or "—", b["received_date"], b["expiry_date"], b["location_name"], b["status"]
            ])
        cls._format_columns_and_borders(ws1, 4)

        # Sheet 2: 25% Barrel Usage History
        ws2 = wb.create_sheet(title="Barrel History")
        cls._apply_excel_styling(ws2, "25% Ammonia Barrel History & Usage")
        headers2 = ["Action Date", "Time", "Barrel ID", "Batch Number", "Original (kg)", "Quantity Used / Transferred", "Remaining (kg)", "Usage Purpose", "Supplier / Customer", "Ref No", "Operator"]
        ws2.append([])
        ws2.append(headers2)
        for h in audit_data.get("barrel_history", []):
            ws2.append([
                h["action_date"], h["action_time"], h["barrel_id"], h["batch_number"],
                h["original_quantity"], h["used_or_transferred_quantity"], h["remaining_quantity"],
                h["usage_purpose"], h.get("supplier_customer_name") or "—", h.get("reference_number") or "—", h["operator_name"]
            ])
        cls._format_columns_and_borders(ws2, 4)

        # Sheet 3: 10% Ammonia Preparation History
        ws3 = wb.create_sheet(title="10% Preparation Log")
        cls._apply_excel_styling(ws3, "10% Ammonia Dilution Log (Fixed Ratio)")
        headers3 = ["Production Date", "Ref No", "Source 25% Barrel", "Source 25% Batch", "25% Used (kg)", "Water Added (kg)", "10% Produced (kg)", "New 10% Batch", "Expiry Date", "Operator", "Remarks"]
        ws3.append([])
        ws3.append(headers3)
        for d in audit_data.get("dilutions", []):
            ws3.append([
                d["production_date"], d["reference_number"], d["source_barrel_id"], d["source_25_batch"],
                d["source_quantity"], d["water_quantity"], d["produced_quantity"],
                d["new_10_batch"], d["expiry_date"], d["operator_name"], d.get("remarks") or "—"
            ])
        cls._format_columns_and_borders(ws3, 4)

        out = io.BytesIO()
        wb.save(out)
        return out.getvalue()

    @classmethod
    def export_monthly_usage_excel(cls, year_month: str, rows: List[Dict[str, Any]]) -> bytes:
        """Generates Excel workbook for Monthly Chemical Usage."""
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = f"Usage {year_month}"
        cls._apply_excel_styling(ws, f"Monthly Chemical Usage ({year_month})")

        headers = ["Chemical Name", "Batch Number", "Expiry Date", "Storage Location", "Transactions Count", "Total Qty Used", "Unit", "Purposes", "Reference Numbers"]
        ws.append([])
        ws.append(headers)

        for r in rows:
            ws.append([
                r["chemical_name"], r["batch_number"], r["expiry_date"], r["storage_location"],
                r["transaction_count"], r["total_quantity_used"], r["unit"], r.get("purposes") or "—", r.get("reference_numbers") or "—"
            ])
        cls._format_columns_and_borders(ws, 4)

        out = io.BytesIO()
        wb.save(out)
        return out.getvalue()

    @classmethod
    def export_physical_vs_system_excel(cls, rows: List[Dict[str, Any]]) -> bytes:
        """Generates Excel workbook for Physical vs System Stock checks."""
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Physical vs System"
        cls._apply_excel_styling(ws, "Physical vs System Stock Reconciliation")

        headers = ["Date", "Time", "Store Location", "Chemical Name", "Physical Qty", "System Qty", "Variance", "Unit", "Inspector Remarks", "Operator"]
        ws.append([])
        ws.append(headers)

        for r in rows:
            ws.append([
                r["date"], r["time"], r["storage_location_name"], r["chemical_name"],
                r["physical_quantity"], r["system_quantity"], r["variance"], r["unit"],
                r.get("remarks") or "—", r["operator_name"]
            ])
        cls._format_columns_and_borders(ws, 4)

        out = io.BytesIO()
        wb.save(out)
        return out.getvalue()

    @classmethod
    def export_historical_records_excel(cls, records: List[Dict[str, Any]]) -> bytes:
        """Generates Excel workbook for Historical Transaction Ledger."""
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Historical Transactions"
        cls._apply_excel_styling(ws, "Historical Transaction Audit Ledger")

        headers = ["Date", "Time", "Type", "Chemical", "Batch Number", "Barrel ID", "Quantity", "Unit", "Purpose / Action", "Supplier / Customer", "Location", "Ref No", "Operator"]
        ws.append([])
        ws.append(headers)

        for r in records:
            ws.append([
                r["trans_date"], r["trans_time"], r["trans_type"], r["chemical_name"],
                r.get("batch_number") or "—", r.get("barrel_id") or "—", r["quantity"],
                r["unit"], r.get("purpose") or "—", r.get("supplier_customer_name") or "—",
                r.get("location_name") or "—", r.get("reference_number") or "—", r.get("operator_name") or "—"
            ])
        cls._format_columns_and_borders(ws, 4)

        out = io.BytesIO()
        wb.save(out)
        return out.getvalue()

    # --- REAL PDF EXPORT GENERATORS VIA REPORTLAB ---
    @classmethod
    def export_report_pdf(cls, title: str, headers: List[str], data_rows: List[List[Any]], landscape_mode: bool = True) -> bytes:
        """Generates a professional corporate PDF report with header, table, and timestamps."""
        buffer = io.BytesIO()
        pagesize = landscape(letter) if landscape_mode else letter
        doc = SimpleDocTemplate(
            buffer,
            pagesize=pagesize,
            leftMargin=30,
            rightMargin=30,
            topMargin=30,
            bottomMargin=30
        )

        styles = getSampleStyleSheet()
        title_style = ParagraphStyle(
            'ReportTitle',
            parent=styles['Heading1'],
            fontSize=16,
            textColor=colors.HexColor('#0C4A6E'),
            spaceAfter=4,
            fontName='Helvetica-Bold'
        )
        subtitle_style = ParagraphStyle(
            'ReportSubtitle',
            parent=styles['Normal'],
            fontSize=9,
            textColor=colors.HexColor('#64748B'),
            spaceAfter=12,
            fontName='Helvetica'
        )
        cell_style = ParagraphStyle(
            'TableCell',
            parent=styles['Normal'],
            fontSize=8,
            textColor=colors.HexColor('#1E293B'),
            fontName='Helvetica'
        )
        header_cell_style = ParagraphStyle(
            'TableHeaderCell',
            parent=styles['Normal'],
            fontSize=8,
            textColor=colors.white,
            fontName='Helvetica-Bold',
            alignment=1 # Center
        )

        elements = []
        elements.append(Paragraph("CHEMICAL STORE MANAGEMENT SYSTEM", title_style))
        elements.append(Paragraph(f"Official Store Document: <b>{title}</b> &bull; Generated: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}", subtitle_style))
        elements.append(Spacer(1, 6))

        # Format table data
        formatted_table_data = []
        formatted_table_data.append([Paragraph(h, header_cell_style) for h in headers])

        for row in data_rows:
            formatted_row = []
            for cell in row:
                val = str(cell) if cell is not None else "—"
                formatted_row.append(Paragraph(val, cell_style))
            formatted_table_data.append(formatted_row)

        table = Table(formatted_table_data, repeatRows=1)
        table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#0284C7')),
            ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
            ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
            ('TOPPADDING', (0, 0), (-1, -1), 4),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F8FAFC')]),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#CBD5E1')),
        ]))

        elements.append(table)
        doc.build(elements)
        return buffer.getvalue()
